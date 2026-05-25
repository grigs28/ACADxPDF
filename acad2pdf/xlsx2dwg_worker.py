# acad2pdf/xlsx2dwg_worker.py
"""xlsx → DWG 转换（全专业），供 Worker 调用。"""

import logging
import os
import shutil
import sys
import time

log = logging.getLogger("acad2pdf")

_TOOLS_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "tools", "xlsx2dwg_net"
)
_DLL_PATH = os.path.join(_TOOLS_DIR, "XlsxToDwg", "bin", "Release", "XlsxToDwg.dll")
_DEFAULT_TEMPLATE = r"C:\opt\ACADxPDF\Template\A1.dwt"

# 序号 → 专业目录名（ZIP 内显示）
SHEET_INDEX_MAP = {
    1: "建筑",
    2: "结构",
    3: "给排水",
    4: "暖通",
    5: "电气",
}


def _ensure_tools():
    if _TOOLS_DIR not in sys.path:
        sys.path.insert(0, _TOOLS_DIR)


def _build_dll():
    import subprocess
    if os.path.isfile(_DLL_PATH) and time.time() - os.path.getmtime(_DLL_PATH) < 3600:
        return _DLL_PATH
    dll_dir = os.path.join(_TOOLS_DIR, "XlsxToDwg")
    r = subprocess.run(
        ["dotnet", "build", "-c", "Release"],
        cwd=dll_dir, capture_output=True, text=True, timeout=60
    )
    if r.returncode != 0:
        log.error("xlsx2dwg DLL build failed: %s", r.stderr)
        return None
    return _DLL_PATH


def _resolve_sheets(xlsx_path, sheets):
    """将序号列表 [1,2,3] 解析为 (sheet_index_0based, dir_name) 列表。"""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    all_names = wb.sheetnames
    wb.close()

    if not sheets:
        return [(i, SHEET_INDEX_MAP.get(i + 1, f"s{i + 1}")) for i in range(len(all_names))]

    result = []
    for s in sheets:
        idx = int(s) - 1  # 1-based → 0-based
        if 0 <= idx < len(all_names):
            dir_name = SHEET_INDEX_MAP.get(idx + 1, f"s{idx + 1}")
            result.append((idx, dir_name))
    return result


def convert_one_xlsx(xlsx_path, output_dir, sheets=None, dll_path=None,
                     template=None, timeout=300):
    """
    转换单个 xlsx 文件（可指定多个 sheet/专业）。

    Args:
        sheets: 1-based 序号列表如 [1,2,3,4,5]，None=全部

    Returns:
        {"ok": bool, "outputs": {dir_name: dwg_path}, "error": str}
    """
    _ensure_tools()
    from xlsx2json import extract_to_json
    from run_xlsx2dwg import _run_one
    import openpyxl

    if not os.path.isfile(xlsx_path):
        return {"ok": False, "outputs": {}, "error": f"File not found: {xlsx_path}"}

    # 解析 sheet 序号 → (0-based index, dir_name)
    resolved = _resolve_sheets(xlsx_path, sheets)
    if not resolved:
        return {"ok": False, "outputs": {}, "error": "No valid sheets"}

    # DLL
    effective_dll = dll_path or _build_dll()
    if not effective_dll or not os.path.isfile(effective_dll):
        return {"ok": False, "outputs": {}, "error": "DLL not available"}

    effective_tpl = template or _DEFAULT_TEMPLATE

    # 读取所有 sheet 名称（用于 extract_to_json 的 sheet_name 参数）
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    all_names = wb.sheetnames
    wb.close()

    outputs = {}
    errors = []
    for idx, dir_name in resolved:
        sheet_name = all_names[idx]
        try:
            sheet_dir = os.path.join(output_dir, dir_name)
            os.makedirs(sheet_dir, exist_ok=True)

            json_path = extract_to_json(xlsx_path, sheet_dir, sheet_name)
            if not json_path:
                errors.append(f"{dir_name}: extract failed")
                continue

            result = _run_one(json_path, effective_dll, sheet_dir,
                              effective_tpl, timeout)
            if result["success"]:
                # 重命名为中文文件名：绿色建筑设计专篇（{专业}）.dwg
                src_dwg = result["dwg"]
                new_name = f"绿色建筑设计专篇（{dir_name}）.dwg"
                dst_dwg = os.path.join(sheet_dir, new_name)
                if os.path.isfile(src_dwg) and src_dwg != dst_dwg:
                    shutil.move(src_dwg, dst_dwg)
                    result["dwg"] = dst_dwg
                outputs[dir_name] = result["dwg"]
                log.info("xlsx2dwg: %s/%s OK (%.1fs)",
                         os.path.basename(xlsx_path), dir_name, result["elapsed"])
            else:
                errors.append(f"{dir_name}: {result['error']}")

            try:
                os.remove(json_path)
            except OSError:
                pass

        except Exception as ex:
            errors.append(f"{dir_name}: {ex}")
            log.error("xlsx2dwg %s/%s error: %s",
                      os.path.basename(xlsx_path), dir_name, ex)

    ok = len(outputs) > 0
    error = "; ".join(errors) if errors else ""
    return {"ok": ok, "outputs": outputs, "error": error}
